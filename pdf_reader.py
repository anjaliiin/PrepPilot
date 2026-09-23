"""
pdf_reader.py - Multi-format Document Text Extraction for PrepPilot
Extracts raw text from PDF, DOCX, and TXT files with robust encoding fallbacks and error handling.
"""

import io
import os
from typing import Union


def extract_text_from_file(file_obj: Union[str, io.BytesIO], filename: str) -> str:
    """
    Extracts plain text from an uploaded file or path (.pdf, .docx, .txt).
    Raises ValueError with descriptive feedback on failure or empty content.
    """
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        return extract_text_from_pdf(file_obj)
    elif ext == ".docx":
        return extract_text_from_docx(file_obj)
    elif ext == ".txt":
        return extract_text_from_txt(file_obj)
    else:
        raise ValueError(f"Unsupported file format '{ext}'. Please upload a .pdf, .docx, or .txt syllabus.")


def extract_text_from_pdf(file_obj: Union[str, io.BytesIO]) -> str:
    """Extracts text from a PDF file using pypdf."""
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfReader
        except ImportError:
            raise ImportError("Neither 'pypdf' nor 'PyPDF2' is installed. Please install 'pypdf'.")

    try:
        reader = PdfReader(file_obj)
        pages_text = []
        for idx, page in enumerate(reader.pages):
            txt = page.extract_text()
            if txt:
                pages_text.append(txt)

        full_text = "\n\n".join(pages_text).strip()
        if not full_text:
            raise ValueError("The uploaded PDF appears to be empty or contains only non-extractable scanned images.")
        return full_text
    except Exception as e:
        if isinstance(e, ValueError):
            raise
        raise ValueError(f"Failed to extract text from PDF: {str(e)}")


def extract_text_from_docx(file_obj: Union[str, io.BytesIO]) -> str:
    """Extracts text from a DOCX file using python-docx."""
    try:
        import docx
    except ImportError:
        raise ImportError("python-docx is not installed. Please install 'python-docx'.")

    try:
        doc = docx.Document(file_obj)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        full_text = "\n".join(paragraphs).strip()
        if not full_text:
            raise ValueError("The uploaded DOCX file appears to have no text content.")
        return full_text
    except Exception as e:
        if isinstance(e, ValueError):
            raise
        raise ValueError(f"Failed to extract text from DOCX: {str(e)}")


def extract_text_from_txt(file_obj: Union[str, io.BytesIO]) -> str:
    """Extracts text from a TXT file with UTF-8 and Latin-1 fallbacks."""
    try:
        if isinstance(file_obj, str):
            with open(file_obj, "rb") as f:
                raw_bytes = f.read()
        elif hasattr(file_obj, "read"):
            raw_bytes = file_obj.read()
            if hasattr(file_obj, "seek"):
                file_obj.seek(0)
        else:
            raw_bytes = bytes(file_obj)

        try:
            full_text = raw_bytes.decode("utf-8")
        except UnicodeDecodeError:
            full_text = raw_bytes.decode("latin-1")

        full_text = full_text.strip()
        if not full_text:
            raise ValueError("The uploaded TXT file is empty.")
        return full_text
    except Exception as e:
        if isinstance(e, ValueError):
            raise
        raise ValueError(f"Failed to read TXT file: {str(e)}")
